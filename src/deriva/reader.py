"""Reading one ENEMDU variable dictionary workbook as INEC publishes it."""

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from deriva import DerivaError

HEADER = "Nombre del campo"
YEAR = re.compile(r"(?:19|20)\d\d")


@dataclass
class Dictionary:
    label: str
    variables: dict[str, str]  # variable name -> description, in sheet order


def round_label(path: Path) -> str:
    """The last 19xx/20xx year in the file name, or the file name without extension."""
    years = YEAR.findall(path.stem)
    return years[-1] if years else path.stem


def distinct_labels(label_a: str, label_b: str) -> tuple[str, str]:
    """Suffix the two labels with _a and _b when they would otherwise be the same."""
    if label_a == label_b:
        return f"{label_a}_a", f"{label_b}_b"
    return label_a, label_b


def read_dictionary(path: str | Path) -> Dictionary:
    """Read the variable names and descriptions below the `Nombre del campo` header."""
    path = Path(path)
    if not path.is_file():
        raise DerivaError(f"{path}: file not found")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise DerivaError(f"{path}: cannot be opened as an .xlsx workbook ({exc})") from exc

    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        variables = _read_variables(path, rows)
    finally:
        workbook.close()
    return Dictionary(label=round_label(path), variables=variables)


def _read_variables(path: Path, rows) -> dict[str, str]:
    header_found = False
    variables: dict[str, str] = {}
    for row_number, row in enumerate(rows, start=1):
        name, description = (tuple(row) + (None, None))[:2]
        if not header_found:
            header_found = _text(name) == HEADER
            continue
        name, description = _text(name), _text(description, strip=False)
        if not name and not description:
            continue
        if not name:
            raise DerivaError(f"{path}: row {row_number} has a description but no variable name")
        if name in variables:
            raise DerivaError(f"{path}: variable {name!r} appears more than once (again in row {row_number})")
        variables[name] = description

    if not header_found:
        raise DerivaError(
            f"{path}: no header row found; looked for {HEADER!r} in the first column"
        )
    return variables


def _text(value, strip: bool = True) -> str:
    if value is None:
        return ""
    text = str(value)
    return text.strip() if strip else text
